#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
SEED do orçado · PASSO 2 — da planilha de ANÁLISE revisada para o arquivo de
import do Planorc (Orçar → Importar Baseline).

  python3 scripts/seed_import.py [analise.xlsx] [import.xlsx]

Lê as 12 colunas VERDES (a proposta, já com as suas edições), ignora as colunas
de realizado e as de apoio, e escreve o layout que a tela espera:

  Empresa | Filial | ItemOrcamento | Centro De Custo | Area | Divisão | BU |
  Histórico | <12 colunas de mês, com header do tipo DATA>

O import acha as colunas pelo NOME e os meses por serem Data — por isso as de
análise (que têm header de texto) não interferem, mas aqui saem de fora de
qualquer forma: o arquivo de import é enxuto de propósito, para o que é enviado
ao banco ser exatamente o que foi revisado, sem coluna a mais para interpretar.

A planilha de análise é aberta com data_only=True: o Excel precisa tê-la salvo
ao menos uma vez para as fórmulas terem valor em cache. Se "Proposto" vier
vazio, foi isso — abra, salve e rode de novo. (O total é recalculado aqui a
partir dos 12 meses, então ele não depende do cache; o aviso vale para quem
conferir a coluna.)
"""
import sys, os, re
from collections import defaultdict
from datetime import datetime
from openpyxl import load_workbook, Workbook

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
_pos = [a for a in sys.argv[1:] if not a.startswith('--')]
ENTRADA = _pos[0] if _pos else os.path.join(RAIZ, 'seed2027_analise.xlsx')
SAIDA   = _pos[1] if len(_pos) > 1 else os.path.join(RAIZ, 'seed2027_import.xlsx')


def main():
    wb = load_workbook(ENTRADA, data_only=True)
    linhas, ano_destino = [], None
    zerados = vazios = 0

    for ws in wb.worksheets:
        hdr = [c.value for c in ws[1]]
        if 'Empresa' not in hdr or 'Conta' not in hdr:
            continue                                   # aba Resumo e afins
        IX = {h: i for i, h in enumerate(hdr) if isinstance(h, str)}
        # as 12 colunas da proposta são as do ANO DESTINO: 'jan/27'...'dez/27'.
        # O ano-base aparece como 'jan/26' — os dois sufixos distinguem.
        anos = sorted({int(m.group(1)) for h in hdr if isinstance(h, str)
                       for m in [re.fullmatch(r'[a-z]{3}/(\d{2})', h)] if m})
        if len(anos) < 2:
            print(f'⚠ aba "{ws.title}" sem as colunas de mês esperadas — ignorada'); continue
        suf = f'{anos[-1]:02d}'
        ano_destino = 2000 + anos[-1]
        meses = [IX[f'{m}/{suf}'] for m in ['jan','fev','mar','abr','mai','jun','jul','ago','set','out','nov','dez']]
        iObs = IX.get('Decisão / observação')

        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[IX['Empresa']]:
                continue
            vals = [float(row[i] or 0) for i in meses]
            if not any(round(v, 2) for v in vals):
                zerados += 1; continue
            fil = str(row[IX['Filial']] or '').strip()
            if not fil:
                vazios += 1
            obs = str(row[iObs]).strip() if iObs is not None and row[iObs] else ''
            linhas.append({
                'empresa': str(row[IX['Empresa']]).strip(),
                'filial':  fil,
                'item':    str(row[IX['Conta']]).strip(),
                'cc':      str(row[IX['CC']] or '').strip(),
                'hist':    obs or f'Seed {ano_destino} · {row[IX["Classe"]]}',
                'vals':    vals,
            })

    if not linhas:
        print('Nada a importar — nenhuma linha com valor.'); sys.exit(1)

    # mesma combinação em duas linhas (edição manual pode duplicar) — somar,
    # senão o import grava as duas e a segunda some na agregação da tela
    agg, hist = defaultdict(lambda: [0.0] * 12), {}
    for L in linhas:
        k = (L['empresa'], L['filial'], L['item'], L['cc'])
        for i in range(12): agg[k][i] += L['vals'][i]
        hist.setdefault(k, L['hist'])

    wb2 = Workbook(); ws2 = wb2.active; ws2.title = 'Dados'
    ws2.append(['Empresa', 'Filial', 'ItemOrcamento', 'Centro De Custo', 'Area', 'Divisão', 'BU', 'Histórico']
               + [datetime(ano_destino, m, 1) for m in range(1, 13)])
    for k in sorted(agg):
        ws2.append([k[0], k[1], k[2], k[3], '', '', '', hist[k]] + [round(v, 2) for v in agg[k]])
    for i in range(2, ws2.max_row + 1):
        for c in range(9, 21): ws2.cell(i, c).number_format = '#,##0.00'
    ws2.column_dimensions['C'].width = 14; ws2.column_dimensions['H'].width = 30
    wb2.save(SAIDA)

    tot = sum(sum(v) for v in agg.values())
    print(f'{len(linhas)} linhas lidas -> {len(agg)} combinações em {SAIDA}')
    print(f'total {ano_destino}: {tot:,.2f}')
    if zerados: print(f'   {zerados} linha(s) zerada(s) na revisão — não entram no import')
    if vazios:  print(f'   ⚠ {vazios} linha(s) SEM FILIAL — entrariam com filial nula')


if __name__ == '__main__':
    main()
